"""Merge result stays saveable even when the on-disk cache cannot be written."""

from __future__ import annotations

import os
import sys
import tempfile
import unittest
from unittest.mock import patch

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from openpyxl import Workbook

import merge_engine


def _csv(folder: str, name: str, body: str) -> str:
    path = os.path.join(folder, name)
    with open(path, "w", encoding="utf-8", newline="") as fh:
        fh.write(body)
    return path


class MergeSaveTests(unittest.TestCase):
    def test_existing_file_name_column_still_saves(self):
        with tempfile.TemporaryDirectory() as folder:
            path = _csv(folder, "cam.csv", "File name,Value\ncamera,1\n")
            res = merge_engine.merge_files([path])
            self.assertTrue(res["ok"], res)
            self.assertEqual(res["rows"], 1)
            status = merge_engine.merge_status()
            self.assertTrue(status["ready"])
            out = os.path.join(folder, "merged.csv")
            saved = merge_engine.save_merged(out)
            self.assertTrue(saved["ok"], saved)
            with open(out, encoding="utf-8-sig") as fh:
                text = fh.read()
            self.assertIn("camera", text)
            self.assertIn("File name (from file)", text)

    def test_save_works_when_pickle_fails(self):
        with tempfile.TemporaryDirectory() as folder:
            path = _csv(folder, "a.csv", "Value\n3\n4\n")
            with patch.object(merge_engine.pd.DataFrame, "to_pickle", side_effect=OSError("disk full")):
                res = merge_engine.merge_files([path])
            self.assertTrue(res["ok"], res)
            self.assertEqual(res["rows"], 2)
            self.assertEqual(res["cache_path"], "")
            out = os.path.join(folder, "merged.csv")
            saved = merge_engine.save_merged(out)
            self.assertTrue(saved["ok"], saved)
            self.assertEqual(saved["rows"], 2)

    def test_data_behind_empty_cover_sheet_is_merged(self):
        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, "vision.xlsx")
            book = Workbook()
            cover = book.active
            cover.title = "Cover"
            cover.append(["Title"])
            cover.append(["Vision export"])
            data = book.create_sheet("Measurements")
            data.append(["DATE", "Value"])
            for i in range(120):
                data.append(["2026-10-0700-11-52", i])
            book.save(path)
            book.close()

            res = merge_engine.merge_files([path])
            self.assertTrue(res["ok"], res)
            self.assertEqual(res["rows"], 120)
            self.assertIn("DATE", res["columns"])
            self.assertNotIn("Sheet", res["columns"])

    def test_sheet_contains_still_selects_one_sheet(self):
        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, "two.xlsx")
            book = Workbook()
            first = book.active
            first.title = "Summary"
            first.append(["Note"])
            first.append(["ignore"])
            second = book.create_sheet("Measurements")
            second.append(["Value"])
            second.append([7])
            second.append([8])
            book.save(path)
            book.close()

            res = merge_engine.merge_files([path], sheet_contains="meas")
            self.assertTrue(res["ok"], res)
            self.assertEqual(res["rows"], 2)
            self.assertEqual(list(merge_engine._MERGED["Value"]), [7, 8])
